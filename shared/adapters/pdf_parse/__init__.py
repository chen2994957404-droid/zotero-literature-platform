# -*- coding: utf-8 -*-
"""pdf_parse · PDF 解析基础件（原子能力：PDF → 结构化文本+图坐标）

职责：把一个 PDF 解析成 full.md（全文）+ layout.json（版面/图坐标）+ images/ + 原PDF副本。
这是「原子模块层」的一块——精读/结构化抽取/向量化三条工作流都依赖它。
底层用 MineRU 云端 API（VLM 模型，处理公式/表格/版面）。

原子模块的特征：只做「PDF→解析结果」这一件不可再分的事，不依赖任何上层模块。

对外接口：
  - parse_pdf(pdf_path, out_dir) → out_dir（含 full.md/layout.json/images/*_origin.pdf）
  - link_origin(out_dir, pdf_path) → 把 *_origin.pdf 换成指向正本的硬链接（省一份空间）
  - same_geometry(a, b)          → 两个 PDF 页面几何是否一致
                                    已解析过（out_dir 有 layout.json）则直接复用，省 MineRU。
  - parse_docx(path, out_dir)    → out_dir（只有 full.md：文字 + 表格，python-docx 读，不花额度）
  - parse_document(path, out_dir) → 按扩展名分派到上面两个。**精读 / 取全文 / 落地流水线
                                    三处都只调这一个**（2026-09-13 收拢，此前三处各写一遍）。
  - parse_pdf_text(pdf_path, out_dir) → 快速文本层：PyMuPDF 本地抽字，几秒出 full.md（无表格结构、无图），
                                    打 `.tier_text` 标记。之后 parse_pdf（MineRU）成功会覆盖它并去掉标记。
  - parse_document_text(path, out_dir) → 快速层的分派：pdf → parse_pdf_text，docx → parse_docx。
  - tier(out_dir)                 → 'structured'（MineRU / docx）/ 'text'（只有快速层）/ 'none'

## 两层（2026-10-04，Claude Science 的需求文档）

MineRU 是云端排队，忙起来一篇 pending 半小时（2026-10-02 两篇都卡在 pending、0 字可读，
PDF 明明已经在盘上）。所以拿到 PDF 先出**快速文本层**（本地、几秒、不联网），
按节按段马上能读；MineRU 在后台补表格与版面，成功了升级成 structured，失败了文本层照旧可用。

配置（环境变量）：
  - MINERU_TOKEN : MineRU API token（必须；无默认，密钥不硬编码）
"""
import os, re, json, time, zipfile, io, urllib.request, urllib.error
import urllib.parse as _up, http.client as _hc

from shared.kernel.log import get_logger

log = get_logger('pdf_parse')

BASE = 'https://mineru.net/api/v4'


class PDFParseError(Exception):
    pass



POLL_EVERY_S = 8          # 多久问一次 MineRU
POLL_MAX_S = 900          # 最多等多久（15 分钟）
PENDING_TTL_S = 48 * 3600  # 超时没等完的任务记多久（过了就当它丢了，重新上传）

def _token():
    """取 MineRU token：走 config 原子模块（环境变量 → .env），避免子进程拿不到。"""
    t = os.environ.get('MINERU_TOKEN')
    if not t:
        try:
            from shared.kernel.config import get_key
            t = get_key('MINERU_TOKEN')
        except Exception:
            t = ''
    if not t:
        raise PDFParseError(
            '未找到 MINERU_TOKEN。请在项目根 .env 写 MINERU_TOKEN=你的token，'
            '或设环境变量后重启进程')
    return t


def _api(path, method='GET', body=None):
    data = json.dumps(body).encode() if body else None
    req = urllib.request.Request(BASE + path, data=data, method=method,
        headers={'Authorization': f'Bearer {_token()}', 'Content-Type': 'application/json'})
    return json.loads(urllib.request.urlopen(req, timeout=60).read())


def is_parsed(out_dir):
    """out_dir 是否已有解析结果（layout.json 存在即视为已解析）。"""
    return os.path.exists(os.path.join(out_dir, 'layout.json'))


def parse_pdf(pdf_path, out_dir, reuse=True):
    """解析 PDF 到 out_dir。已解析过且 reuse=True 则直接复用（省 MineRU 额度）。

    产出：out_dir/{full.md, layout.json, *_origin.pdf, images/}。返回 out_dir。
    先用 vlm 档；MineRU 那边回「parsing failed」就换 pipeline 档再试一次
    （2026-09-18：两份 6–9 MB 的 Nature Mater. / Sci. Adv. 正文 vlm 连败三次，文件本身没坏）。
    """
    os.makedirs(out_dir, exist_ok=True)
    if reuse and is_parsed(out_dir):
        return out_dir
    try:
        return _parse_once(pdf_path, out_dir, model_version='vlm', ocr=True)
    except PDFParseError as e:
        if '解析失败' not in str(e):
            raise
        return _parse_once(pdf_path, out_dir, model_version='pipeline', ocr=False)


_PENDING = '.mineru_pending.json'


def _pending_get(out_dir, pdf_path, model_version):
    """上次超时时还在 MineRU 那边排队的任务 → batch_id，没有 / 过期 / 文件变了 → ''。"""
    p = os.path.join(out_dir, _PENDING)
    try:
        with io.open(p, encoding='utf-8') as f:
            d = json.load(f)
    except (OSError, ValueError):
        return ''
    if (d.get('model') == model_version and d.get('size') == os.path.getsize(pdf_path)
            and time.time() - d.get('at', 0) < PENDING_TTL_S):
        return d.get('batch_id') or ''
    return ''


def _pending_set(out_dir, pdf_path, model_version, batch_id):
    try:
        with io.open(os.path.join(out_dir, _PENDING), 'w', encoding='utf-8') as f:
            json.dump({'batch_id': batch_id, 'model': model_version, 'size': os.path.getsize(pdf_path),
                       'at': time.time()}, f)
    except OSError:
        pass


def _pending_clear(out_dir):
    try:
        os.remove(os.path.join(out_dir, _PENDING))
    except OSError:
        pass


def _parse_once(pdf_path, out_dir, model_version, ocr):
    fname = os.path.basename(pdf_path)
    # 0. 上次超时、MineRU 那边还在排着的任务：**接着查它，不重新上传**（2026-10-02）。
    #    重新上传 = 排到队尾重排 —— MineRU 忙时一篇 Adv. Mater. 排了 17 分钟还是 pending，每次超时重试都从头排。
    batch_id = _pending_get(out_dir, pdf_path, model_version)
    if batch_id:
        log.info(f'{fname} 接着查上次没等完的 MineRU 任务 {batch_id}（不重新上传、不重新排队）')
        return _poll_and_fetch(pdf_path, out_dir, model_version, ocr, batch_id)
    # 1. 申请上传地址
    r = _api('/file-urls/batch', 'POST', {
        "enable_formula": True, "enable_table": True, "language": "en", "model_version": model_version,
        "files": [{"name": fname, "is_ocr": ocr, "data_id": "zot_" + str(int(time.time()))}]})
    batch_id = r['data']['batch_id']
    upload_url = r['data']['file_urls'][0]

    # 2. PUT 上传（无 Content-Type 头，避免 OSS 签名不匹配 —— 踩坑 #1）
    with open(pdf_path, 'rb') as f:
        pdf_bytes = f.read()
    u = _up.urlparse(upload_url)
    conn = _hc.HTTPSConnection(u.netloc, timeout=180)
    conn.request('PUT', u.path + '?' + u.query, body=pdf_bytes, headers={})
    resp = conn.getresponse(); resp.read()
    status = resp.status; conn.close()
    if status not in (200, 201):
        raise PDFParseError(f'上传失败 HTTP {status}')
    _pending_set(out_dir, pdf_path, model_version, batch_id)     # 万一下面等超时，下次接着查这个任务
    return _poll_and_fetch(pdf_path, out_dir, model_version, ocr, batch_id)


def _poll_and_fetch(pdf_path, out_dir, model_version, ocr, batch_id):

    # 3. 轮询（字段结构见踩坑 #2）。最多等 POLL_MAX_S：原来 40×8 秒≈5 分钟，大文献（12 MB 的 Adv. Mater.）
    # 在 MineRU 那边排队 + 解析常超过它，平台先放弃了、MineRU 其实还在跑（2026-10-02）
    zip_url, st = None, '?'
    for _ in range(POLL_MAX_S // POLL_EVERY_S):
        time.sleep(POLL_EVERY_S)
        r = _api(f'/extract-results/batch/{batch_id}')
        res = r['data']['extract_result'][0]
        st = res['state']
        if st == 'done':
            zip_url = res['full_zip_url']; break
        if st == 'failed':
            # 把 MineRU 回的整条记录和 batch_id 留下：它的 err_msg 只有一句通用话，
            # 拿 batch_id 才能去它后台查（2026-09-18 一晚 10 次失败，事后只剩这一句）
            log.warn('MineRU 解析失败 batch=%s model=%s ocr=%s 返回=%s' % (batch_id, model_version, ocr, res))
            _pending_clear(out_dir)                  # 明确失败了：下次重新上传（换档也走这条）
            raise PDFParseError('解析失败: ' + res.get('err_msg', '') + f'（batch {batch_id}，{model_version}）')
    if not zip_url:
        log.warn('MineRU 解析等了 %d 秒还没好 batch=%s 最后状态=%s' % (POLL_MAX_S, batch_id, st))
        raise PDFParseError(f'解析超时（等了 {POLL_MAX_S // 60} 分钟，MineRU 状态 {st}，batch {batch_id}；'
                            f'任务还在 MineRU 那边排着，再提交这篇会接着等它，不会重新排队）')

    # 4. 下载解压
    zip_bytes = urllib.request.urlopen(zip_url, timeout=120).read()
    zipfile.ZipFile(io.BytesIO(zip_bytes)).extractall(out_dir)
    _pending_clear(out_dir)
    _tier_text_clear(out_dir)                 # MineRU 的 full.md 已经覆盖了快速层
    link_origin(out_dir, pdf_path)
    return out_dir


def same_geometry(a, b):
    """两个 PDF 页数、每页宽高与旋转是否全部一致（裁图坐标能否通用的判据）。没装 PyMuPDF 返回 False。"""
    try:
        import fitz
        da, db = fitz.open(a), fitz.open(b)
        if len(da) != len(db):
            return False
        return all((pa.rect.width, pa.rect.height, pa.rotation) == (pb.rect.width, pb.rect.height, pb.rotation)
                   for pa, pb in zip(da, db))
    except Exception:
        return False


def link_origin(out_dir, pdf_path):
    """MineRU 的 `*_origin.pdf` 就是输入 PDF 的原样副本 —— 换成指向正本的硬链接，不再占第二份空间。

    2026-09-15 量过主力机：1424 个 origin.pdf 共 9.2 GB。MineRU 会把 PDF 重新保存一遍，
    字节不完全一样，但**页数与每页尺寸全部一致**（968/968），裁图坐标对着正本裁出来的图
    与对着 origin 裁的**像素级相同**（25 篇实测最大平均差 0.00/255）—— 所以换成硬链接是安全的。
    硬链接在 NTFS 上不要管理员权限、同一份数据两个名字，裁图（figure_crop）照旧按名字找得到。
    只在页面几何一致时换；换不成（跨盘、权限、没装 PyMuPDF）就保留副本。返回换了几个。
    """
    if not (pdf_path and os.path.isfile(pdf_path)):
        return 0
    n = 0
    for f in os.listdir(out_dir):
        if not f.endswith('_origin.pdf'):
            continue
        dup = os.path.join(out_dir, f)
        try:
            if os.path.samefile(dup, pdf_path) or not same_geometry(dup, pdf_path):
                continue
            tmp = dup + '.lnk'
            os.link(pdf_path, tmp)
            os.replace(tmp, dup)
            n += 1
        except OSError:
            continue
    return n


def parse_docx(path, out_dir, reuse=True):
    """.docx → out_dir/full.md（文字 + 表格，表格行拼成 `a | b | c`）。不联网、不花额度。

    从 tools/deepread/si.py 下沉（2026-09-13）：落地流水线也要读 docx 的 SI，
    两个使用者 + 用了第三方库 python-docx → 按规矩住 adapters。
    表格务必取：SI 的投料量、配比常常只在表里。
    """
    os.makedirs(out_dir, exist_ok=True)
    md = os.path.join(out_dir, 'full.md')
    if reuse and os.path.exists(md):
        return out_dir
    try:
        import docx
    except ImportError:
        raise PDFParseError('需要 python-docx：pip install python-docx')
    doc = docx.Document(path)
    parts = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
    for tb in doc.tables:
        for row in tb.rows:
            cells = [c.text.strip() for c in row.cells]
            if any(cells):
                parts.append(' | '.join(cells))
    with io.open(md, 'w', encoding='utf-8') as fh:
        fh.write('\n\n'.join(parts))
    return out_dir


def is_text_parsed(out_dir):
    """out_dir 里有没有 full.md（docx 解析没有 layout.json，只能看这个）。"""
    return os.path.exists(os.path.join(out_dir, 'full.md'))


def real_ext(path_or_bytes, fallback=''):
    """看文件头判真实类型：'.pdf' / '.docx' / fallback。

    2026-09-20 主力机 ingest 里 12 篇 SI 反复报 PackageNotFoundError：Elsevier / Nature 给的
    `mmc1.docx` 链接下来的其实是 PDF（文件头 %PDF-1.7），按扩展名当 docx 读当然炸。
    落地和解析都别信扩展名，信文件头。docx 的头是 zip 的 PK\x03\x04。
    """
    if isinstance(path_or_bytes, (bytes, bytearray)):
        head = bytes(path_or_bytes[:8])
    elif path_or_bytes:
        try:
            with io.open(str(path_or_bytes), 'rb') as fh:
                head = fh.read(8)
        except OSError:
            return fallback
    else:
        head = b''
    if head.startswith(b'%PDF'):
        return '.pdf'
    if head.startswith(b'PK\x03\x04'):
        return '.docx'
    return fallback


def parse_document(path, out_dir, reuse=True):
    """PDF 或 docx → out_dir/full.md。**按文件头分派**，扩展名只是兜底；两样都不认抛 PDFParseError。"""
    ext = real_ext(path, os.path.splitext(path or '')[1].lower())
    if ext == '.pdf':
        return parse_pdf(path, out_dir, reuse=reuse)
    if ext == '.docx':
        return parse_docx(path, out_dir, reuse=reuse)
    raise PDFParseError(f'不会解析这种文件：{ext or "(无扩展名)"}（只认 .pdf / .docx）')


# ══════════════════════════════════════════════════════════════════════
# 快速文本层（PyMuPDF，本地几秒）
# ══════════════════════════════════════════════════════════════════════

TIER_TEXT, TIER_STRUCTURED, TIER_NONE = 'text', 'structured', 'none'
_TIER_MARK = '.tier_text'
MIN_TEXT_CHARS = 1500     # 整篇抽出来不到这么多字 = 没有文字层（扫描件），快速层帮不上，等 MineRU 的 OCR


def tier(out_dir):
    """这个解析目录到哪一档：有 full.md 且没有快速层标记 = structured（MineRU 或 docx 直读）。"""
    if not os.path.exists(os.path.join(out_dir, 'full.md')):
        return TIER_NONE
    return TIER_TEXT if os.path.exists(os.path.join(out_dir, _TIER_MARK)) else TIER_STRUCTURED


def _tier_text_clear(out_dir):
    try:
        os.remove(os.path.join(out_dir, _TIER_MARK))
    except OSError:
        pass


_SECTION_WORDS = re.compile(
    r'(?i)^(?:\d+(?:\.\d+)*\.?\s+|[IVX]+\.\s+)?(?:abstract|introduction|background|results?|discussion|'
    r'results and discussion|conclusions?|summary|experimental(?: section)?|methods?|materials and methods|'
    r'materials|characterization|acknowledge?ments?|references|supporting information|'
    r'associated content|author information|notes|conflicts? of interest|data availability)\b')
_NUMBERED_HEAD = re.compile(r'^\d+(?:\.\d+){0,3}\.?\s+[A-Z]')
_MULTI_NUM_HEAD = re.compile(r'^\d+\.\d+(?:\.\d+){0,2}\.?\s+[A-Z][A-Za-z]')
_CAPTION_START = re.compile(r'(?i)^(?:fig(?:ure)?|table|scheme)\.?\s*S?\d+')


def _norm_repeat(t):
    """页眉页脚判重用：数字抹成 #（页码、卷期每页不同，其余相同）。"""
    return re.sub(r'\d+', '#', t.strip().lower())[:80]


def _blocks(doc):
    """→ [(页号, 文字, 字号, 是否加粗, 相对纵坐标, 开头的加粗/斜体段)]，每个文字块一条，行内拼好、断词接上。

    「开头的加粗/斜体段」用来认**接排标题**：`2.2. Preparation of PDBS. The mixture was…` 这种
    标题和正文挤在同一块里（Elsevier 常见），整块看不像标题，只有开头那截换了字体。
    """
    out = []
    for pno, page in enumerate(doc):
        h = page.rect.height or 1
        for b in page.get_text('dict').get('blocks', []):
            if b.get('type') != 0:
                continue
            lines, sizes, bold, total, lead, lead_open = [], [], 0, 0, '', True
            for ln in b.get('lines', []):
                spans = [s for s in ln.get('spans', []) if s.get('text', '').strip()]
                if not spans:
                    continue
                lines.append(''.join(s['text'] for s in ln['spans']).strip())
                for s in spans:
                    n = len(s['text'].strip())
                    sizes.append((s.get('size', 0), n))
                    total += n
                    font = (s.get('font') or '').lower()
                    is_b = bool(s.get('flags', 0) & 16) or 'bold' in font
                    styled = is_b or bool(s.get('flags', 0) & 2) or 'italic' in font   # Elsevier 二级标题是斜体
                    if is_b:
                        bold += n
                    if lead_open:
                        if styled:
                            lead += s['text']
                        else:
                            lead_open = False
            if not lines:
                continue
            text = ''
            for ln in lines:                   # 行尾连字符 + 下一行小写开头 = 一个词被折断了
                if text.endswith('-') and ln[:1].islower():
                    text = text[:-1] + ln
                else:
                    text = (text + ' ' + ln) if text else ln
            size = max(sizes, key=lambda x: x[1])[0] if sizes else 0
            out.append((pno, text, round(size * 2) / 2, bool(total) and bold / total > 0.6,
                        b['bbox'][1] / h, ' '.join(lead.split())))
    return out


_BULLET = re.compile(r'^[■▪●•◆▶⬛]+\s*')   # ACS 的 ■ INTRODUCTION 之类


def _is_heading(t, size, body, bold):
    """一整块像不像标题。"""
    t0 = _BULLET.sub('', t).strip()
    if len(t0) >= 140 or t0.endswith(('.', ',', ';')) or _CAPTION_START.match(t0):
        return False
    # 图里的矢量字（子图标号 a / b / c、坐标轴、单位）也是一个个文字块：字母太少的一律不算
    if (len(re.findall(r'[A-Za-z]', t0)) < 4 or not re.search(r'[A-Za-z]{3}', t0)
            or t0.lower().startswith(('doi', 'http', 'www.'))):
        return False
    if size >= body + 1.5:
        return True
    if _SECTION_WORDS.fullmatch(t0.rstrip(':')) and len(t0) < 60:
        return True
    if bold and size >= body - 0.6:
        # 加粗的短行：章节词 / 编号开头的一律算；别的要像一句标题（≥2 个词、大写开头、不太长）
        if _SECTION_WORDS.match(t0) or _NUMBERED_HEAD.match(t0):
            return True
        return len(t0) < 100 and len(t0.split()) >= 2 and t0[:1].isupper()
    # 多级编号（2.1 / 3.2.1）本身就说明是小节标题，不管字体
    return bool(_MULTI_NUM_HEAD.match(t0)) and len(t0) < 120


def _run_in_head(t, lead):
    """接排标题：开头加粗段像编号标题 / 章节词，且后面还有正文 → (标题, 正文)，否则 None。"""
    lead = _BULLET.sub('', lead.strip()).rstrip(':').strip()
    t = _BULLET.sub('', t.strip())
    if not lead or len(lead) > 120 or len(lead) >= len(t) - 20 or not t.startswith(lead[:10]):
        return None
    if not (_NUMBERED_HEAD.match(lead) or _SECTION_WORDS.match(lead)):
        return None
    return lead.rstrip('.'), t[len(lead):].lstrip(' .:')


def text_markdown(doc):
    """一个打开的 PyMuPDF 文档 → Markdown 文本（标题按字号 / 加粗 / 章节词认，正文一块一段）。

    只求「按节按段能读」，不求版面：表格会散成几行字，图只剩图注 —— 那是 MineRU 那一档的事。
    """
    blocks = _blocks(doc)
    if not blocks:
        return ''
    # 正文字号 = 按字数加权最多的那个
    weight = {}
    for b in blocks:
        weight[b[2]] = weight.get(b[2], 0) + len(b[1])
    body = max(weight, key=lambda k: weight[k])
    # 页眉页脚：贴着页顶 / 页底、在三分之一以上的页上重复出现
    pages = len(doc)
    seen = {}
    for p, t, _s, _b, y, _l in blocks:
        if (y < 0.08 or y > 0.92) and len(t) < 160:
            seen.setdefault(_norm_repeat(t), set()).add(p)
    repeat = {k for k, ps in seen.items() if pages >= 3 and len(ps) >= max(3, pages // 3)}

    parts, title_done = [], False
    for p, t, s, bold, y, lead in blocks:
        if (y < 0.08 or y > 0.92) and (_norm_repeat(t) in repeat or re.fullmatch(r'\d{1,4}', t.strip())):
            continue
        if not title_done and p == 0 and s >= body + 4 and len(t) < 300:
            parts.append('# ' + t)
            title_done = True
            continue
        if _is_heading(t, s, body, bold):
            parts.append('## ' + _BULLET.sub('', t).strip())
            continue
        ri = _run_in_head(t, lead)
        if ri:
            parts.append('## ' + ri[0])
            if ri[1]:
                parts.append(ri[1])
            continue
        parts.append(t)
    return '\n\n'.join(parts) + '\n'


def parse_pdf_text(pdf_path, out_dir, reuse=True):
    """PDF → out_dir/full.md（快速文本层）。本地 PyMuPDF，不联网、不花额度，一篇几秒。

    已有 full.md（任何一档）且 reuse=True 时不动它 —— 不拿快速层去盖 MineRU 的结果。
    没有文字层（扫描件）抛 PDFParseError：快速层帮不上，只能等 MineRU 的 OCR。
    """
    os.makedirs(out_dir, exist_ok=True)
    md_path = os.path.join(out_dir, 'full.md')
    if reuse and os.path.exists(md_path):
        return out_dir
    try:
        import fitz
    except ImportError:
        raise PDFParseError('需要 PyMuPDF：pip install PyMuPDF')
    try:
        doc = fitz.open(pdf_path)
    except Exception as e:
        raise PDFParseError(f'PDF 打不开：{type(e).__name__}: {str(e)[:80]}')
    try:
        md = text_markdown(doc)
        pages = len(doc)
    finally:
        doc.close()
    if len(md) < MIN_TEXT_CHARS:
        raise PDFParseError(f'PDF 几乎没有文字层（{pages} 页只抽到 {len(md)} 字，多半是扫描件），要等 MineRU 的 OCR')
    with io.open(os.path.join(out_dir, _TIER_MARK), 'w', encoding='utf-8') as fh:
        json.dump({'engine': 'pymupdf', 'at': time.strftime('%Y-%m-%d %H:%M:%S'),
                   'pages': pages, 'chars': len(md)}, fh)
    tmp = md_path + '.tmp'
    with io.open(tmp, 'w', encoding='utf-8') as fh:
        fh.write(md)
    os.replace(tmp, md_path)
    return out_dir


def parse_document_text(path, out_dir, reuse=True):
    """快速层的分派：pdf → parse_pdf_text；docx 本来就是直读文字 + 表格（就是它的最终档）。"""
    ext = real_ext(path, os.path.splitext(path or '')[1].lower())
    if ext == '.pdf':
        return parse_pdf_text(path, out_dir, reuse=reuse)
    if ext == '.docx':
        return parse_docx(path, out_dir, reuse=reuse)
    raise PDFParseError(f'不会解析这种文件：{ext or "(无扩展名)"}（只认 .pdf / .docx）')


def check_token(timeout=20):
    """MineRU token 还有效吗？返回 (ok, 说明)。**零成本，不产生解析任务**。

    做法：拿一个不存在的 batch id 去查结果 ——
      · token 有效 → HTTP 200 + `task not found or expire`（业务层说找不到）
      · token 无效 → HTTP 401 `user authenticate failed`
    实测确认过两种响应（2026-08-28）。这是目前找到的唯一免费校验方式：
    MineRU 没有「查账号/查额度」这类接口。
    """
    try:
        tok = _token()
    except PDFParseError as e:
        return False, str(e)[:60]
    req = urllib.request.Request(
        BASE + '/extract-results/batch/zzzznotexist',
        headers={'Authorization': 'Bearer ' + tok})
    try:
        urllib.request.urlopen(req, timeout=timeout)
        return True, '有效'
    except urllib.error.HTTPError as e:
        if e.code == 401:
            return False, 'token 无效或已过期，去 mineru.net 重新申请'
        return None, f'查不了：HTTP {e.code}'
    except Exception as e:
        return None, f'连不上 MineRU：{type(e).__name__}'
